"""
Tests para el sistema de importación de Excel.
Verifica el funcionamiento de los importadores y endpoints.
"""
import os
import tempfile
from django.test import TestCase
from django.contrib.auth import get_user_model
from django.urls import reverse
from rest_framework.test import APITestCase
from rest_framework import status
from openpyxl import Workbook

from authentication.models import Role
from entities.models import Proveedor, Cliente, Unidad
from inventory.models import GPS, SIMCard
from sales.models import Ventas
from .importers import GPSImporter, SIMCardImporter, ClienteImporter, ProveedorImporter, VentasImporter

User = get_user_model()


class BaseImporterTestCase(TestCase):
    """Test case base para los importadores."""
    
    def setUp(self):
        """Configuración inicial para los tests."""
        # Crear rol de administrador
        self.admin_role = Role.objects.create(
            nombre='administrador',
            descripcion='Administrador del sistema'
        )
        
        # Crear proveedor de prueba
        self.proveedor = Proveedor.objects.create(
            nombre='Proveedor Test',
            ruc='20123456789',
            direccion='Dirección Test',
            contacto='Contacto Test',
            celular='987654321'
        )
    
    def create_excel_file(self, data, headers):
        """Crear archivo Excel temporal para testing."""
        wb = Workbook()
        ws = wb.active
        
        # Agregar headers
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        
        # Agregar datos
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # Guardar en archivo temporal
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        wb.save(temp_file.name)
        temp_file.close()
        
        return temp_file.name


class GPSImporterTestCase(BaseImporterTestCase):
    """Tests para GPSImporter."""
    
    def test_import_gps_success(self):
        """Test importación exitosa de GPS."""
        headers = ['fecha_compra', 'imei', 'marca', 'modelo', 'numero_factura', 'proveedor_ruc']
        data = [
            ['2024-01-15', '123456789012345', 'Teltonika', 'FMB920', 'F001-123', '20123456789']
        ]
        
        excel_file = self.create_excel_file(data, headers)
        
        try:
            importer = GPSImporter()
            result = importer.import_from_excel(excel_file)
            
            self.assertEqual(result['success_count'], 1)
            self.assertEqual(result['error_count'], 0)
            self.assertEqual(GPS.objects.count(), 1)
            
            gps = GPS.objects.first()
            self.assertEqual(gps.imei, '123456789012345')
            self.assertEqual(gps.marca, 'Teltonika')
            
        finally:
            os.unlink(excel_file)
    
    def test_import_gps_invalid_imei(self):
        """Test importación con IMEI inválido."""
        headers = ['fecha_compra', 'imei', 'marca', 'modelo', 'numero_factura', 'proveedor_ruc']
        data = [
            ['2024-01-15', '123', 'Teltonika', 'FMB920', 'F001-123', '20123456789']
        ]
        
        excel_file = self.create_excel_file(data, headers)
        
        try:
            importer = GPSImporter()
            result = importer.import_from_excel(excel_file)
            
            self.assertEqual(result['success_count'], 0)
            self.assertEqual(result['error_count'], 1)
            self.assertEqual(GPS.objects.count(), 0)
            
        finally:
            os.unlink(excel_file)


class ClienteImporterTestCase(BaseImporterTestCase):
    """Tests para ClienteImporter."""
    
    def test_import_cliente_success(self):
        """Test importación exitosa de cliente."""
        headers = ['nombre', 'ruc', 'direccion', 'contacto', 'celular', 'correo']
        data = [
            ['Transportes ABC', '20987654321', 'Av. Principal 123', 'Juan Pérez', '987654321', 'test@example.com']
        ]
        
        excel_file = self.create_excel_file(data, headers)
        
        try:
            importer = ClienteImporter()
            result = importer.import_from_excel(excel_file)
            
            self.assertEqual(result['success_count'], 1)
            self.assertEqual(result['error_count'], 0)
            self.assertEqual(Cliente.objects.count(), 1)
            
            cliente = Cliente.objects.first()
            self.assertEqual(cliente.nombre, 'Transportes ABC')
            self.assertEqual(cliente.ruc, '20987654321')
            
        finally:
            os.unlink(excel_file)


class ImportAPITestCase(APITestCase):
    """Tests para los endpoints de importación."""
    
    def setUp(self):
        """Configuración inicial para los tests de API."""
        # Crear rol de administrador
        self.admin_role = Role.objects.create(
            nombre='administrador',
            descripcion='Administrador del sistema'
        )
        
        # Crear usuario administrador
        self.admin_user = User.objects.create_user(
            username='admin',
            email='admin@test.com',
            password='testpass123',
            rol=self.admin_role
        )
        
        # Crear usuario normal
        self.normal_role = Role.objects.create(
            nombre='usuario',
            descripcion='Usuario normal'
        )
        self.normal_user = User.objects.create_user(
            username='user',
            email='user@test.com',
            password='testpass123',
            rol=self.normal_role
        )
        
        # Crear proveedor de prueba
        self.proveedor = Proveedor.objects.create(
            nombre='Proveedor Test',
            ruc='20123456789',
            direccion='Dirección Test',
            contacto='Contacto Test',
            celular='987654321'
        )
    
    def test_import_excel_admin_access(self):
        """Test que solo administradores pueden importar."""
        self.client.force_authenticate(user=self.admin_user)
        
        # Crear archivo Excel de prueba
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value='nombre')
        ws.cell(row=2, column=1, value='Test Cliente')
        
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        wb.save(temp_file.name)
        temp_file.close()
        
        try:
            with open(temp_file.name, 'rb') as f:
                response = self.client.post(
                    reverse('imports:import_excel'),
                    {'file': f, 'module': 'cliente'},
                    format='multipart'
                )
            
            # Debería permitir el acceso
            self.assertNotEqual(response.status_code, status.HTTP_403_FORBIDDEN)
            
        finally:
            os.unlink(temp_file.name)
    
    def test_import_excel_normal_user_denied(self):
        """Test que usuarios normales no pueden importar."""
        self.client.force_authenticate(user=self.normal_user)
        
        response = self.client.post(
            reverse('imports:import_excel'),
            {'module': 'cliente'},
            format='multipart'
        )
        
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
    
    def test_import_template_endpoint(self):
        """Test endpoint de plantillas."""
        self.client.force_authenticate(user=self.admin_user)
        
        response = self.client.get(reverse('imports:import_template', kwargs={'module': 'gps'}))
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('required_columns', response.data)
        self.assertIn('example_data', response.data)
    
    def test_import_status_endpoint(self):
        """Test endpoint de estado."""
        self.client.force_authenticate(user=self.admin_user)
        
        response = self.client.get(reverse('imports:import_status'))
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('available_modules', response.data)
        self.assertIn('supported_formats', response.data)
        
        # Verificar que el endpoint devuelve información de todos los módulos
        self.assertIn('gps', response.data['modules'])
        self.assertIn('simcard', response.data['modules'])
        self.assertIn('cliente', response.data['modules'])
        self.assertIn('proveedor', response.data['modules'])
        self.assertIn('ventas', response.data['modules'])


class VentasImporterTestCase(TestCase):
    """Tests para VentasImporter"""
    
    def setUp(self):
        """Configuración inicial para los tests"""
        # Crear cliente
        self.cliente = Cliente.objects.create(
            nombre='Cliente Test',
            ruc='20123456789',
            direccion='Av. Test 123',
            contacto='Juan Test',
            celular='987654321',
            correo='test@cliente.com'
        )
        
        # Crear unidad
        self.unidad = Unidad.objects.create(
            placa='ABC-123',
            marca='Toyota',
            modelo='Hilux',
            año=2020,
            cliente=self.cliente
        )
        
        self.importer = VentasImporter()
    
    def create_test_excel(self, data):
        """Crea un archivo Excel temporal para testing"""
        df = pd.DataFrame(data)
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            df.to_excel(tmp.name, index=False)
            return tmp.name
    
    def test_import_ventas_success(self):
        """Test de importación exitosa de ventas"""
        data = [{
            'mes': '2024-01-01',
            'fecha_pago': '14:30:00',
            'numero_operacion': 'OP123456',
            'tipo_pago': 'transferencia',
            'banco': 'BCP',
            'numero_factura': 'F001-123',
            'fecha_generacion_factura': '2024-01-15 14:30:00',
            'cliente_ruc': '20123456789',
            'descripcion': 'Servicio GPS mensual',
            'unidad_placa': 'ABC-123',
            'precio': 118.00,
            'estado': 'pendiente'
        }]
        
        file_path = self.create_test_excel(data)
        
        try:
            result = self.importer.import_from_excel(file_path)
            
            # Verificar resultado
            self.assertEqual(result['success_count'], 1)
            self.assertEqual(result['error_count'], 0)
            self.assertEqual(result['total_rows'], 1)
            
            # Verificar que se creó la venta
            venta = Ventas.objects.get(numero_factura='F001-123')
            self.assertEqual(venta.cliente, self.cliente)
            self.assertEqual(venta.unidad, self.unidad)
            self.assertEqual(venta.precio, 118.00)
            self.assertEqual(venta.tipo_pago, 'transferencia')
            
            # Verificar cálculos automáticos
            self.assertAlmostEqual(float(venta.importe), 100.00, places=2)
            self.assertAlmostEqual(float(venta.igv), 18.00, places=2)
            self.assertAlmostEqual(float(venta.total), 118.00, places=2)
            
        finally:
            os.unlink(file_path)
    
    def test_import_ventas_invalid_ruc(self):
        """Test con RUC de cliente inválido"""
        data = [{
            'mes': '2024-01-01',
            'fecha_pago': '14:30:00',
            'numero_operacion': 'OP123456',
            'tipo_pago': 'transferencia',
            'banco': 'BCP',
            'numero_factura': 'F001-124',
            'fecha_generacion_factura': '2024-01-15 14:30:00',
            'cliente_ruc': '123',  # RUC inválido
            'descripcion': 'Servicio GPS mensual',
            'unidad_placa': 'ABC-123',
            'precio': 118.00
        }]
        
        file_path = self.create_test_excel(data)
        
        try:
            result = self.importer.import_from_excel(file_path)
            
            # Verificar que hay errores
            self.assertEqual(result['success_count'], 0)
            self.assertEqual(result['error_count'], 1)
            self.assertIn('RUC del cliente debe tener 11 dígitos', str(result['errors']))
            
        finally:
            os.unlink(file_path)
    
    def test_import_ventas_cliente_not_exists(self):
        """Test con cliente que no existe"""
        data = [{
            'mes': '2024-01-01',
            'fecha_pago': '14:30:00',
            'numero_operacion': 'OP123456',
            'tipo_pago': 'transferencia',
            'banco': 'BCP',
            'numero_factura': 'F001-125',
            'fecha_generacion_factura': '2024-01-15 14:30:00',
            'cliente_ruc': '20999999999',  # Cliente que no existe
            'descripcion': 'Servicio GPS mensual',
            'unidad_placa': 'ABC-123',
            'precio': 118.00
        }]
        
        file_path = self.create_test_excel(data)
        
        try:
            result = self.importer.import_from_excel(file_path)
            
            # Verificar que hay errores
            self.assertEqual(result['success_count'], 0)
            self.assertEqual(result['error_count'], 1)
            self.assertIn('Cliente con RUC 20999999999 no existe', str(result['errors']))
            
        finally:
            os.unlink(file_path)
    
    def test_import_ventas_unidad_not_belongs_to_client(self):
        """Test con unidad que no pertenece al cliente"""
        # Crear otro cliente
        otro_cliente = Cliente.objects.create(
            nombre='Otro Cliente',
            ruc='20987654321',
            direccion='Av. Otro 456',
            contacto='Pedro Otro',
            celular='912345678',
            correo='otro@cliente.com'
        )
        
        data = [{
            'mes': '2024-01-01',
            'fecha_pago': '14:30:00',
            'numero_operacion': 'OP123456',
            'tipo_pago': 'transferencia',
            'banco': 'BCP',
            'numero_factura': 'F001-126',
            'fecha_generacion_factura': '2024-01-15 14:30:00',
            'cliente_ruc': '20987654321',  # Otro cliente
            'descripcion': 'Servicio GPS mensual',
            'unidad_placa': 'ABC-123',  # Unidad del primer cliente
            'precio': 118.00
        }]
        
        file_path = self.create_test_excel(data)
        
        try:
            result = self.importer.import_from_excel(file_path)
            
            # Verificar que hay errores
            self.assertEqual(result['success_count'], 0)
            self.assertEqual(result['error_count'], 1)
            self.assertIn('no pertenece al cliente', str(result['errors']))
            
        finally:
            os.unlink(file_path)
    
    def test_import_ventas_duplicate_factura(self):
        """Test con número de factura duplicado"""
        # Crear primera venta
        Ventas.objects.create(
            mes='2024-01-01',
            fecha_pago='10:00:00',
            numero_operacion='OP000001',
            tipo_pago='efectivo',
            banco='BCP',
            numero_factura='F001-999',
            fecha_generacion_factura='2024-01-01 10:00:00',
            cliente=self.cliente,
            descripcion='Venta existente',
            unidad=self.unidad,
            precio=100.00
        )
        
        data = [{
            'mes': '2024-01-01',
            'fecha_pago': '14:30:00',
            'numero_operacion': 'OP123456',
            'tipo_pago': 'transferencia',
            'banco': 'BCP',
            'numero_factura': 'F001-999',  # Factura duplicada
            'fecha_generacion_factura': '2024-01-15 14:30:00',
            'cliente_ruc': '20123456789',
            'descripcion': 'Servicio GPS mensual',
            'unidad_placa': 'ABC-123',
            'precio': 118.00
        }]
        
        file_path = self.create_test_excel(data)
        
        try:
            result = self.importer.import_from_excel(file_path)
            
            # Verificar que hay errores
            self.assertEqual(result['success_count'], 0)
            self.assertEqual(result['error_count'], 1)
            self.assertIn('Factura F001-999 ya existe', str(result['errors']))
            
        finally:
            os.unlink(file_path)
    
    def test_import_ventas_invalid_tipo_pago(self):
        """Test con tipo de pago inválido"""
        data = [{
            'mes': '2024-01-01',
            'fecha_pago': '14:30:00',
            'numero_operacion': 'OP123456',
            'tipo_pago': 'bitcoin',  # Tipo de pago inválido
            'banco': 'BCP',
            'numero_factura': 'F001-127',
            'fecha_generacion_factura': '2024-01-15 14:30:00',
            'cliente_ruc': '20123456789',
            'descripcion': 'Servicio GPS mensual',
            'unidad_placa': 'ABC-123',
            'precio': 118.00
        }]
        
        file_path = self.create_test_excel(data)
        
        try:
            result = self.importer.import_from_excel(file_path)
            
            # Verificar que hay errores
            self.assertEqual(result['success_count'], 0)
            self.assertEqual(result['error_count'], 1)
            self.assertIn('Tipo de pago inválido', str(result['errors']))
            
        finally:
            os.unlink(file_path)
